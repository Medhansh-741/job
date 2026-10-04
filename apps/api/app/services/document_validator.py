import io
import os
import re
import uuid
import zipfile
from dataclasses import dataclass
from typing import Tuple

import docx
import pdfplumber
from fastapi import HTTPException, status

MAX_FILE_SIZE = 5 * 1024 * 1024  # 5 MB
MAX_UNCOMPRESSED_DOCX_SIZE = 25 * 1024 * 1024  # 25 MB max decompressed
MAX_DOCX_ENTRIES = 200
MIN_READABLE_CHARS = 50

PDF_MAGIC = b"%PDF-"
ZIP_MAGIC = b"PK\x03\x04"


class DocumentValidationError(HTTPException):
    def __init__(self, code: str, message: str):
        super().__init__(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail={"code": code, "message": message},
        )


@dataclass
class ValidatedDocument:
    raw_bytes: bytes
    extracted_text: str
    extension: str
    safe_filename: str
    content_type: str


def validate_filename(filename: str) -> str:
    """Validate filename for path traversal, extension allowlist, and double extensions."""
    if not filename:
        raise DocumentValidationError(
            "INVALID_FILENAME", "Filename cannot be empty."
        )

    # Basic normalization
    clean_name = os.path.basename(filename).strip()

    # Reject path traversal attempts
    if ".." in clean_name or "/" in clean_name or "\\" in clean_name:
        raise DocumentValidationError(
            "INVALID_FILENAME", "Invalid characters or path traversal in filename."
        )

    parts = clean_name.split(".")
    if len(parts) < 2:
        raise DocumentValidationError(
            "INVALID_EXTENSION", "File lacks an extension. Only .pdf and .docx are supported."
        )

    ext = f".{parts[-1].lower()}"

    # Check for executable or script double extensions (e.g. resume.pdf.exe, file.docx.sh)
    executable_exts = {
        "exe", "bat", "cmd", "sh", "ps1", "vbs", "js", "ts", "py", "php", "asp",
        "aspx", "jsp", "jar", "bin", "elf", "scr", "pif", "hta", "cpl", "dll"
    }
    for part in parts[1:]:
        if part.lower() in executable_exts:
            raise DocumentValidationError(
                "INVALID_FILENAME",
                "Disallowed file extension detected. Executable and script files are strictly blocked.",
            )

    if ext not in [".pdf", ".docx"]:
        raise DocumentValidationError(
            "INVALID_EXTENSION",
            f"Unsupported file extension '{ext}'. Only .pdf and .docx files are permitted.",
        )

    return ext


def validate_magic_bytes(content: bytes, ext: str):
    """Verify true file format using magic bytes."""
    if len(content) < 8:
        raise DocumentValidationError(
            "FILE_EMPTY", "Uploaded file is too small to be a valid document."
        )

    if ext == ".pdf":
        if not content.startswith(PDF_MAGIC):
            raise DocumentValidationError(
                "SPOOFED_FILE_TYPE",
                "File content does not match its .pdf extension (invalid PDF signature).",
            )
    elif ext == ".docx":
        if not content.startswith(ZIP_MAGIC):
            raise DocumentValidationError(
                "SPOOFED_FILE_TYPE",
                "File content does not match its .docx extension (invalid Office OpenXML signature).",
            )


def validate_docx_structure(content: bytes):
    """Deep inspection of DOCX archive: zip bombs, macros, and real Word structure."""
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            infos = z.infolist()

            if len(infos) > MAX_DOCX_ENTRIES:
                raise DocumentValidationError(
                    "DECOMPRESSION_LIMIT_EXCEEDED",
                    "Word document archive contains too many internal files.",
                )

            total_uncompressed = sum(info.file_size for info in infos)
            if total_uncompressed > MAX_UNCOMPRESSED_DOCX_SIZE:
                raise DocumentValidationError(
                    "DECOMPRESSION_BOMB",
                    "Document exceeds maximum safe decompression limit (possible zip bomb).",
                )

            namelist = z.namelist()

            # Reject embedded macros
            for name in namelist:
                lower = name.lower()
                if "vbaproject.bin" in lower or lower.endswith(".vba") or "macros" in lower:
                    raise DocumentValidationError(
                        "MACRO_NOT_ALLOWED",
                        "Documents containing embedded macros (VBA) are blocked for security.",
                    )

            # Must contain actual Word document body
            if "word/document.xml" not in namelist:
                raise DocumentValidationError(
                    "INVALID_DOCX_FORMAT",
                    "The uploaded file is a generic ZIP archive, not a valid Word document.",
                )
    except zipfile.BadZipFile:
        raise DocumentValidationError(
            "CORRUPT_DOCUMENT", "The Word document archive is corrupted and cannot be read."
        )


def extract_and_verify_text(content: bytes, ext: str) -> str:
    """Extract text from PDF or DOCX and verify it is not empty or blank."""
    extracted_text = ""

    if ext == ".pdf":
        try:
            with pdfplumber.open(io.BytesIO(content)) as pdf:
                if len(pdf.pages) == 0:
                    raise DocumentValidationError(
                        "CORRUPT_DOCUMENT", "PDF document contains 0 pages."
                    )
                pages_text = []
                for p in pdf.pages:
                    txt = p.extract_text()
                    if txt:
                        pages_text.append(txt)
                extracted_text = "\n".join(pages_text).strip()
        except DocumentValidationError:
            raise
        except Exception as e:
            raise DocumentValidationError(
                "CORRUPT_DOCUMENT", f"Failed to parse PDF document: {str(e)}"
            )
    elif ext == ".docx":
        try:
            doc = docx.Document(io.BytesIO(content))
            paragraphs = [p.text for p in doc.paragraphs if p.text]
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text:
                            paragraphs.append(cell.text)
            extracted_text = "\n".join(paragraphs).strip()
        except DocumentValidationError:
            raise
        except Exception as e:
            raise DocumentValidationError(
                "CORRUPT_DOCUMENT", f"Failed to parse Word document: {str(e)}"
            )

    # Check for empty / unreadable / blank content
    alnum_count = len(re.findall(r"[a-zA-Z0-9]", extracted_text))
    if alnum_count < MIN_READABLE_CHARS:
        raise DocumentValidationError(
            "NO_READABLE_TEXT",
            "No readable text found in this resume. Please ensure the document is not blank or an un-scanned image.",
        )

    return extracted_text


def validate_document(filename: str, content: bytes) -> ValidatedDocument:
    """Run complete 5-layer validation pipeline on uploaded document."""
    # 1. Size Envelope
    size = len(content)
    if size == 0:
        raise DocumentValidationError(
            "FILE_EMPTY", "Uploaded file is empty (0 bytes)."
        )
    if size > MAX_FILE_SIZE:
        raise DocumentValidationError(
            "FILE_TOO_LARGE",
            f"File size ({size / (1024*1024):.1f} MB) exceeds maximum limit of 5MB.",
        )

    # 2. Filename & Extension
    ext = validate_filename(filename)

    # 3. Magic Bytes
    validate_magic_bytes(content, ext)

    # 4. Deep Structure Inspection
    if ext == ".docx":
        validate_docx_structure(content)

    # 5. Content Extraction & Blankness Check
    extracted_text = extract_and_verify_text(content, ext)

    # Generate safe random UUID filename
    safe_filename = f"{uuid.uuid4()}{ext}"
    content_type = (
        "application/pdf"
        if ext == ".pdf"
        else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

    return ValidatedDocument(
        raw_bytes=content,
        extracted_text=extracted_text,
        extension=ext,
        safe_filename=safe_filename,
        content_type=content_type,
    )
